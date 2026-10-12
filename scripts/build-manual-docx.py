from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "docs" / "MailPost-manual.docx"

INK = RGBColor(0x1C, 0x19, 0x14)
MIST = RGBColor(0x5C, 0x51, 0x44)
AMBER = RGBColor(0x8A, 0x5A, 0x12)


def set_run_font(run, name="微软雅黑", size=11, bold=False, color=INK):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = name
    r = run._element.get_or_add_rPr()
    rFonts = r.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run_font(run, size=18 if level == 1 else 14, bold=True, color=INK if level == 1 else AMBER)
    return p


def add_body(doc, text, bold_parts=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


def add_mixed(doc, parts):
    """parts: list of (text, bold)"""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    for text, bold in parts:
        run = p.add_run(text)
        set_run_font(run, size=11, bold=bold)
    return p


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(0.75)
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=11, bold=True)
        r2 = p.add_run(text)
        set_run_font(r2, size=11)
    else:
        run = p.add_run(text)
        set_run_font(run, size=11)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(0.75)
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shading = hdr[i]._tePr if False else hdr[i]._tc.get_or_add_tcPr()
        from docx.oxml import parse_xml

        shading.append(parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="2A251E"/>'))
    for r_i, row in enumerate(rows):
        cells = table.rows[r_i + 1].cells
        for c_i, val in enumerate(row):
            cells[c_i].text = ""
            p = cells[c_i].paragraphs[0]
            run = p.add_run(val)
            set_run_font(run, size=10)
    doc.add_paragraph()


def main():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.4)
        section.right_margin = Cm(2.4)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    t = title.add_run("MailPost 操作手册")
    set_run_font(t, size=26, bold=True)
    sub = doc.add_paragraph()
    s = sub.add_run("Gmail 发信工作台  ·  接入邮箱  ·  养号  ·  写信")
    set_run_font(s, size=12, color=MIST)
    note = doc.add_paragraph()
    n = note.add_run("打开 http://localhost:3000   启动命令：在项目目录运行 npm run dev")
    set_run_font(n, size=10, color=MIST)

    add_heading(doc, "系统能做什么", 1)
    add_body(doc, "本系统只用 Gmail SMTP 发信。侧栏六个入口如下。数据存在 data/mailpost.json，真实邮件走 Gmail。")
    add_table(
        doc,
        ["菜单", "做什么"],
        [
            ["写信", "选 Gmail、填收件人、发一封"],
            ["Gmail", "接入 / 更新发件邮箱"],
            ["联系人", "自己添加收件人"],
            ["话术", "保存、修改可复用的标题和正文"],
            ["已发送", "查看发出记录"],
            ["设置", "签名名称、地址、退订用的站点地址"],
        ],
    )

    add_heading(doc, "一、添加发件邮箱（Gmail）", 1)
    add_mixed(
        doc,
        [
            ("系统只接受 ", False),
            ("@gmail.com", True),
            ("。密码必须是 ", False),
            ("应用专用密码", True),
            ("，不是网页登录密码。", False),
        ],
    )

    add_heading(doc, "1. 在 Google 账号里准备 SMTP", 2)
    add_number(doc, "用要发信的 Gmail 登录。")
    add_number(doc, "打开 https://myaccount.google.com/security ，开启「两步验证」。")
    add_number(doc, "打开 https://myaccount.google.com/apppasswords 。")
    add_number(doc, "应用选「邮件」，设备选「Windows 计算机」，生成。")
    add_number(doc, "得到 16 位密码（中间可能有空格，填进系统时有没有空格都可以）。")
    add_body(doc, "没有两步验证，就看不到应用专用密码页面。")

    add_heading(doc, "2. 在 MailPost 里接入", 2)
    add_number(doc, "打开 Gmail 页：http://localhost:3000/mailboxes")
    add_number(doc, "右侧填写三项：")
    add_bullet(doc, "you@gmail.com", "Gmail 地址：")
    add_bullet(doc, "收件人看到的发件人名字，例如 lily", "显示名：")
    add_bullet(doc, "上一步那 16 位", "应用专用密码：")
    add_number(doc, "点「接入」。")
    add_body(doc, "系统会自动写死 SMTP 主机 smtp.gmail.com、端口 587、STARTTLS，并强制 IPv4（本机开代理时避免连接被掐）。")
    add_body(doc, "接入成功后，左侧出现这只邮箱，状态一般是「养号 D1」。")

    add_heading(doc, "3. 以后改密码", 2)
    add_body(doc, "同一页卡片里「更新应用专用密码」，填新的 16 位，点「更新」。Google 撤销旧应用密码后必须走这一步，否则发信会失败。")

    add_heading(doc, "4. 接入失败常见原因", 2)
    add_table(
        doc,
        ["现象", "原因"],
        [
            ["提示只支持 Gmail", "地址不是 @gmail.com"],
            ["提示已经接入", "同一地址加过了"],
            ["发送时报 Invalid login", "填了登录密码，或应用专用密码失效"],
            [
                "发送时报 Connection closed unexpectedly",
                "代理 / IPv6；系统已按 IPv4 处理，仍失败就检查代理是否放行 smtp.gmail.com:587",
            ],
        ],
    )

    add_heading(doc, "二、养号", 1)
    add_body(doc, "接入那一刻开始计「养号第 1 天」。")
    add_mixed(
        doc,
        [
            ("本系统的养号不是自动互发、自动回信。它只做一件事：", False),
            ("按天数卡死每天能发多少封", True),
            ("，避免新号一天打出去被 Gmail 封 SMTP。", False),
        ],
    )
    add_body(doc, "真正把号养热，要靠你每天用这些额度发正常、不像群发的信，并尽量有人回复。")

    add_heading(doc, "1. 每日额度（系统强制）", 2)
    add_body(doc, "按接入后天数计算，单邮箱上限 40，再高也不会超过 50：")
    add_table(
        doc,
        ["天数", "每天最多"],
        [
            ["第 1–3 天", "5 封"],
            ["第 4–7 天", "10 封"],
            ["第 8–14 天", "20 封"],
            ["第 15–21 天", "30 封"],
            ["第 22 天起", "40 封"],
        ],
    )
    add_body(doc, "额度在 Gmail 页和写信页都会显示，例如「今日 2/5」。用完当天不能再发，第二天自动重置。")
    add_body(doc, "Gmail 页上的「养号 D7」表示现在是第 7 天；变成「就绪」表示已过 21 天爬坡。")

    add_heading(doc, "2. 建议怎么养（你来发，系统来限）", 2)
    add_mixed(doc, [("前 3 天（每天不超过 5 封）：", True)])
    add_bullet(doc, "发给会回你的人：自己的另一个邮箱、同事、熟人")
    add_bullet(doc, "标题、正文每封都不一样，写成对话，不要同一模板连发")
    add_bullet(doc, "让对方回一封，你再回，像正常通信")
    add_bullet(doc, "不要同一封「MailPost test」复制给很多人")
    add_mixed(doc, [("第 4–14 天：", True)])
    add_bullet(doc, "可以慢慢加真实联系人，仍保持个性化")
    add_bullet(doc, "每天尽量在额度内发完，但不要为了发满而群发套话")
    add_bullet(doc, "看「已发送」，有退信就停，先查地址")
    add_mixed(doc, [("第 15 天以后：", True)])
    add_bullet(doc, "额度放开，仍避免短时间对很多 Gmail 发同一套话")
    add_bullet(doc, "进 Promotions 不等于没送到；要进 Primary，内容得像私人信")

    add_heading(doc, "3. 养号期间不要做的事", 2)
    add_bullet(doc, "一天用多台工具、多个 IP 狂发")
    add_bullet(doc, "买名单群发")
    add_bullet(doc, "带促销链接、图片、附件、「限时免费」")
    add_bullet(doc, "忽略退信和「垃圾邮件」投诉")
    add_body(doc, "Gmail 个人账号整体发送量本来就低，适合测试和小规模触达，不适合一天几百封。")

    add_heading(doc, "三、设置签名（建议接入后就填）", 1)
    add_body(doc, "打开设置：http://localhost:3000/compliance")
    add_bullet(doc, "出现在信末签名", "显示名称：")
    add_bullet(doc, "可选，会附在信末", "地址：")
    add_bullet(doc, "用来拼退订链接，默认 http://localhost:3000", "站点：")
    add_body(doc, "每封信末尾会自动带签名和退订地址。这能满足合规，也可能让 Gmail 把信分到 Promotions。一对一若总进促销栏，多半是同一模板连发，或带了退订头。")

    add_heading(doc, "四、联系人", 1)
    add_body(doc, "打开联系人：http://localhost:3000/contacts")
    add_bullet(doc, "只能自己添加，没有系统预置名单")
    add_bullet(doc, "填邮箱、名、姓、公司，点「添加」")
    add_bullet(doc, "可删除")
    add_bullet(doc, "写信页可以从下拉框选用，自动带出邮箱和称呼")
    add_body(doc, "发信时如果地址不在名单里，系统会按你手填的地址发出，并记入联系人。")

    add_heading(doc, "五、话术", 1)
    add_body(doc, "打开话术：http://localhost:3000/sequences")
    add_bullet(doc, "写名称、标题、正文，点「保存」", "新建：")
    add_bullet(doc, "改完再保存，就是更新")
    add_bullet(doc, "可加多封（第 2 封、第 3 封），也可删某一封或整条话术")
    add_bullet(doc, "{{firstName}} {{company}} {{title}} {{industry}} {{lastName}}，发送时按当前收件人替换", "变量（可选）：")
    add_body(doc, "写信页也可以：先「套用话术」选一条，标题正文被填上；改完填话术名称，点「存为话术」或「更新话术」。")

    add_heading(doc, "六、写信发出", 1)
    add_body(doc, "打开写信：http://localhost:3000")
    add_number(doc, "发件 Gmail（已接入的号）")
    add_number(doc, "选用联系人，或手填「发给谁」")
    add_number(doc, "称呼（会替换话术里的 {{firstName}}）")
    add_number(doc, "可选：套用话术")
    add_number(doc, "确认标题、正文")
    add_number(doc, "点「发送」")
    add_body(doc, "当天额度用完按钮会不可点。结果看页面提示，或到「已发送」。对方没看到时：查垃圾箱和 Promotions。第一封常进 Primary，同一模板第二封常进 Promotions。")

    add_heading(doc, "七、已发送", 1)
    add_body(doc, "打开 http://localhost:3000/activity 。这里只有实际发出去的记录。可把某封标成硬退或已回复，硬退会把该联系人抑制，避免再发。")

    add_heading(doc, "八、推荐日常顺序", 1)
    add_number(doc, "Google 开通两步验证 + 应用专用密码")
    add_number(doc, "Gmail 页接入")
    add_number(doc, "设置里填显示名")
    add_number(doc, "前 3 天每天用写信功能发几封给会回信的人（不超过 5 封）")
    add_number(doc, "联系人里加真实收件人")
    add_number(doc, "话术里写自己的正文")
    add_number(doc, "写信套用话术后改一改再发，不要原样群发")

    add_heading(doc, "九、和 Gmail 页签的关系", 1)
    add_bullet(doc, "Gmail 认为是私人来信", "Primary：")
    add_bullet(doc, "Gmail 认为是批量 / 营销", "Promotions：")
    add_body(doc, "系统无法指定进哪个页签。减少进促销栏的做法：每封都不一样、像对话、慢慢发、让对方把你加到通讯录或点「不是促销」。")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
